#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HARTREE_TO_KCAL_MOL = 627.509474
KCAL_PER_MOL_PER_VOLT = 23.06
ABSOLUTE_SHE = 4.28
SHE_TO_SCE = 0.241

# Generic file matcher:
# - must start with S or O
# - followed by one or more digits
# - somewhere later must contain OPT or SPC
# - extension must be .out
FILE_ROLE_RE = re.compile(
    r"^(?P<state>[SO])(?P<id>\d+).*?(?P<calc>OPT|SPC).*?\.out$",
    re.IGNORECASE,
)

# Exact bare FINAL SINGLE POINT ENERGY line only.
# This excludes:
#   FINAL SINGLE POINT ENERGY (Solute-gas-phase)
#   FINAL SINGLE POINT ENERGY (Solute-CPCM)
#   FINAL SINGLE POINT ENERGY (Solvent-CPCM)
FINAL_SPE_RE = re.compile(
    r"^\s*FINAL SINGLE POINT ENERGY\s+(-?\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?)\s*$"
)

# Capture the value immediately followed by "Eh"
EH_VALUE_RE = re.compile(
    r"(-?\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?)\s+Eh\b"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Calculate oxidation potentials from ORCA OPT/SPC output files and save results to Excel."
    )
    parser.add_argument(
        "--directory",
        default=".",
        help="Directory containing ORCA .out files (default: current directory)",
    )
    parser.add_argument(
        "--output",
        default="oxidation_potential_results.xlsx",
        help="Output Excel file (default: oxidation_potential_results.xlsx)",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Print extracted values while processing",
    )
    return parser.parse_args()


def extract_last_eh_value_from_labeled_lines(file_path: Path, label: str) -> float | None:
    """
    Return the last value in Hartree (Eh) found on lines containing the chosen label.
    This avoids accidentally reading the kcal/mol number that is often printed on
    the same line in ORCA output.
    """
    last_value = None

    with file_path.open("r", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            if label in line:
                match = EH_VALUE_RE.search(line)
                if match:
                    last_value = float(match.group(1))

    return last_value


def extract_g_correction(file_path: Path) -> float | None:
    return extract_last_eh_value_from_labeled_lines(file_path, "G-E(el)")


def extract_dgsolv(file_path: Path) -> float | None:
    return extract_last_eh_value_from_labeled_lines(
        file_path, "Free energy of solvation (dGsolv)"
    )


def extract_single_point_energy(file_path: Path) -> float | None:
    """
    Return the last EXACT bare 'FINAL SINGLE POINT ENERGY <value>' in the file.
    Lines with parentheses after the label are intentionally ignored.
    """
    last_value = None

    with file_path.open("r", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            match = FINAL_SPE_RE.match(line)
            if match:
                last_value = float(match.group(1))

    return last_value


def calculate_voltages(delta_g_kcal_mol: float) -> tuple[float, float]:
    """
    Oxidation potential from free energy difference:
        E_vs_SHE = deltaG / 23.06 - 4.28
        E_vs_SCE = E_vs_SHE - 0.241
    """
    v_vs_she = (delta_g_kcal_mol / KCAL_PER_MOL_PER_VOLT) - ABSOLUTE_SHE
    v_vs_sce = v_vs_she - SHE_TO_SCE
    return round(v_vs_she, 2), round(v_vs_sce, 2)


def classify_files(directory: Path) -> tuple[dict[str, dict[str, Path]], list[str]]:
    """
    Build a mapping like:
      pairs['001']['S_OPT'] = Path(...)
      pairs['001']['S_SPC'] = Path(...)
      pairs['001']['O_OPT'] = Path(...)
      pairs['001']['O_SPC'] = Path(...)

    Also return warnings for ambiguous duplicates.
    """
    pairs: dict[str, dict[str, Path]] = defaultdict(dict)
    warnings: list[str] = []

    for file_path in sorted(directory.iterdir()):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() != ".out":
            continue

        match = FILE_ROLE_RE.match(file_path.name)
        if not match:
            continue

        state = match.group("state").upper()
        mol_id = match.group("id")   # keep original zero-padded digits, e.g. "001"
        calc = match.group("calc").upper()
        key = f"{state}_{calc}"

        if key in pairs[mol_id]:
            warnings.append(
                f"Ambiguous duplicate for ID S{mol_id} role {key}: "
                f"keeping {pairs[mol_id][key].name}, ignoring {file_path.name}"
            )
            continue

        pairs[mol_id][key] = file_path

    return dict(pairs), warnings


def autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        max_length = 0
        column = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        ws.column_dimensions[column].width = min(max_length + 2, 40)


def style_worksheet(ws) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions

    autosize_columns(ws)


def process_directory(directory: Path, output_file: Path, debug: bool = False) -> None:
    pairs, warnings = classify_files(directory)

    wb = Workbook()
    ws = wb.active
    ws.title = "Potentials"

    headers = [
        "ID",
        "Status",
        "V vs SHE (V)",
        "V vs SCE (V)",
        "deltaG (kcal/mol)",
        "G_s (Hartree)",
        "G_o (Hartree)",
        "G-E(el) S (Hartree)",
        "SPE S (Hartree)",
        "dGsolv S (Hartree)",
        "G-E(el) O (Hartree)",
        "SPE O (Hartree)",
        "dGsolv O (Hartree)",
        "S OPT file",
        "S SPC file",
        "O OPT file",
        "O SPC file",
        "Notes",
    ]
    ws.append(headers)

    if not pairs:
        ws.append(["", "No matching ORCA .out files were found."])
    else:
        for mol_id in sorted(pairs, key=lambda x: int(x)):
            group = pairs[mol_id]
            display_id = f"S{mol_id}"

            required_keys = ["S_OPT", "S_SPC", "O_OPT", "O_SPC"]
            missing = [key for key in required_keys if key not in group]

            if missing:
                available = ", ".join(sorted(p.name for p in group.values()))
                ws.append([
                    display_id,
                    "MISSING FILES",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    group["S_OPT"].name if "S_OPT" in group else "",
                    group["S_SPC"].name if "S_SPC" in group else "",
                    group["O_OPT"].name if "O_OPT" in group else "",
                    group["O_SPC"].name if "O_SPC" in group else "",
                    f"Missing roles: {', '.join(missing)}. Available files: {available}",
                ])
                continue

            s_opt = group["S_OPT"]
            s_spc = group["S_SPC"]
            o_opt = group["O_OPT"]
            o_spc = group["O_SPC"]

            try:
                g_s = extract_g_correction(s_opt)
                e_s = extract_single_point_energy(s_spc)
                dgsolv_s = extract_dgsolv(s_spc)

                g_o = extract_g_correction(o_opt)
                e_o = extract_single_point_energy(o_spc)
                dgsolv_o = extract_dgsolv(o_spc)

                extracted = {
                    "G-E(el) S": g_s,
                    "SPE S": e_s,
                    "dGsolv S": dgsolv_s,
                    "G-E(el) O": g_o,
                    "SPE O": e_o,
                    "dGsolv O": dgsolv_o,
                }

                missing_values = [name for name, value in extracted.items() if value is None]
                if missing_values:
                    raise ValueError("Missing extracted values: " + ", ".join(missing_values))

                G_s = e_s + dgsolv_s + g_s
                G_o = e_o + dgsolv_o + g_o

                delta_g = (G_o - G_s) * HARTREE_TO_KCAL_MOL
                v_vs_she, v_vs_sce = calculate_voltages(delta_g)

                if debug:
                    print(f"Processing ID {display_id}")
                    print(f"  S OPT:  {s_opt.name}")
                    print(f"  S SPC:  {s_spc.name}")
                    print(f"  O OPT:  {o_opt.name}")
                    print(f"  O SPC:  {o_spc.name}")
                    for name, value in extracted.items():
                        print(f"  {name}: {value}")
                    print(f"  G_s: {G_s:.10f} Eh")
                    print(f"  G_o: {G_o:.10f} Eh")
                    print(f"  deltaG: {delta_g:.2f} kcal/mol")
                    print(f"  V vs SHE: {v_vs_she:.2f} V")
                    print(f"  V vs SCE: {v_vs_sce:.2f} V")

                ws.append([
                    display_id,
                    "OK",
                    v_vs_she,
                    v_vs_sce,
                    round(delta_g, 2),
                    G_s,
                    G_o,
                    g_s,
                    e_s,
                    dgsolv_s,
                    g_o,
                    e_o,
                    dgsolv_o,
                    s_opt.name,
                    s_spc.name,
                    o_opt.name,
                    o_spc.name,
                    "",
                ])

            except Exception as exc:
                ws.append([
                    display_id,
                    "ERROR",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    s_opt.name,
                    s_spc.name,
                    o_opt.name,
                    o_spc.name,
                    str(exc),
                ])

    # Number formats
    for row in ws.iter_rows(min_row=2):
        for idx in [3, 4]:  # V vs SHE, V vs SCE
            if row[idx - 1].value is not None:
                row[idx - 1].number_format = "0.00"
        if row[4].value is not None:  # deltaG
            row[4].number_format = "0.00"
        for idx in range(6, 14):  # Hartree values
            if row[idx - 1].value is not None:
                row[idx - 1].number_format = "0.0000000000"

    style_worksheet(ws)

    if warnings:
        ws_warn = wb.create_sheet("Warnings")
        ws_warn.append(["Type", "Message"])
        for warning in warnings:
            ws_warn.append(["Warning", warning])
        style_worksheet(ws_warn)

    wb.save(output_file)
    print(f"Results saved in {output_file}")


def main() -> None:
    args = parse_args()
    directory = Path(args.directory).expanduser().resolve()
    output_file = Path(args.output).expanduser().resolve()

    if not directory.exists():
        raise SystemExit(f"ERROR: directory does not exist: {directory}")
    if not directory.is_dir():
        raise SystemExit(f"ERROR: path is not a directory: {directory}")

    process_directory(directory, output_file, debug=args.debug)


if __name__ == "__main__":
    main()